"""
apps/documentos/tests.py
Testes minimos dos models de documentos (models basicos: str, unicidade, FK)
e do backend da Fatia 2 (services de renderizacao/PDF e views AJAX).
"""
import io
import html
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date, timedelta
from decimal import Decimal
from unittest import skipUnless
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.management.base import CommandError
from django.conf import settings
from django.contrib.messages import get_messages
from django.db import IntegrityError, connection, transaction
from django.test import SimpleTestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from django_tenants.postgresql_backend.base import FakeTenant
from django_tenants.test.cases import TenantTestCase
from django_tenants.utils import get_public_schema_name, schema_context
from celery.exceptions import SoftTimeLimitExceeded
from docx import Document
from jinja2.exceptions import SecurityError

from apps.contratos.models import Contrato
from apps.imoveis.models import Imovel, Proprietario
from apps.inquilinos.models import Inquilino
from apps.tenants.models import Plano, Tenant

from .models import ContratoDocumentoGerado, ModeloDocumento, VariavelDocumento
from .services import (
    ConversaoDocxErro,
    ModeloDocxInvalido,
    _fmt_money,
    analisar_docx,
    construir_contexto,
    converter_docx_para_pdf,
    MODELOS_PADRAO_DOCX,
    caminho_modelo_padrao_docx,
    criar_documentos_padrao,
    criar_modelos_padrao_docx,
    criar_variaveis_padrao,
    renderizar_docx,
    renderizar_modelo,
    salvar_documento_gerado,
)
from .tasks import ERRO_GENERICO, gerar_documento_docx
from .validators import ValidarDocx, validar_extensao_docx, validar_tamanho_modelo


class ModeloDocumentoTests(TenantTestCase):

    def test_str_usa_tipo_e_titulo(self):
        modelo = ModeloDocumento.objects.create(
            titulo='Contrato Residencial Padrão', tipo='contrato',
        )

        self.assertEqual(str(modelo), 'Contrato de Locação — Contrato Residencial Padrão')


class VariavelDocumentoTests(TenantTestCase):

    def test_slug_deve_ser_unico(self):
        VariavelDocumento.objects.create(
            slug='inquilino.nome', label='Nome do Inquilino', categoria='inquilino',
        )

        with self.assertRaises(IntegrityError):
            VariavelDocumento.objects.create(
                slug='inquilino.nome', label='Duplicado', categoria='inquilino',
            )


class ContratoDocumentoGeradoTests(TenantTestCase):

    def setUp(self):
        self.imovel = Imovel.objects.create(
            codigo='IM-0001', tipo='apartamento', cep='60000000',
            logradouro='Rua Teste', numero='100', bairro='Centro',
            cidade='Fortaleza', estado='CE',
        )
        self.inquilino = Inquilino.objects.create(
            tipo='pf', nome='Rodrigo Oliveira', cpf='02738306006',
            telefone='85999999999', email='pagador@email.com',
            logradouro='Rua Doutor Vargas', numero='150',
            cidade='Porto Alegre', estado='RS', cep='91250000',
        )
        self.contrato = Contrato.objects.create(
            imovel=self.imovel, inquilino=self.inquilino, numero='CT-0001',
            data_inicio=date(2026, 1, 10), data_fim=date(2026, 12, 10),
            dia_vencimento=10, valor_aluguel=Decimal('1500.00'),
        )

    def test_criar_documento_gerado_vinculado_ao_contrato(self):
        documento = ContratoDocumentoGerado.objects.create(
            contrato=self.contrato, titulo='Contrato CT-0001',
            conteudo_final_html='<p>teste</p>',
        )

        self.assertEqual(documento.contrato, self.contrato)
        self.assertEqual(documento.status, 'gerado')
        self.assertEqual(str(documento), f'Contrato CT-0001 — Contrato {self.contrato.pk}')


class RenderizacaoModeloTests(TenantTestCase):

    def setUp(self):
        self.proprietario = Proprietario.objects.create(
            nome='Maria Souza', cpf_cnpj='11122233344',
        )
        self.imovel = Imovel.objects.create(
            codigo='IM-0001', tipo='apartamento', cep='60000000',
            logradouro='Rua Teste', numero='100', bairro='Centro',
            cidade='Fortaleza', estado='CE',
            proprietario=self.proprietario,
        )
        self.inquilino = Inquilino.objects.create(
            tipo='pf', nome='Rodrigo Oliveira', cpf='02738306006',
            telefone='85999999999', email='pagador@email.com',
            logradouro='Rua Doutor Vargas', numero='150',
            cidade='Porto Alegre', estado='RS', cep='91250000',
        )
        self.contrato = Contrato.objects.create(
            imovel=self.imovel, inquilino=self.inquilino, numero='CT-0001',
            data_inicio=date(2026, 1, 10), data_fim=date(2026, 12, 10),
            dia_vencimento=10, valor_aluguel=Decimal('1500.00'),
        )

    def test_construir_contexto_retorna_chaves_esperadas(self):
        contexto = construir_contexto(self.contrato)

        self.assertEqual(contexto['inquilino']['nome'], 'Rodrigo Oliveira')
        self.assertEqual(contexto['imovel']['proprietario_nome'], 'Maria Souza')
        self.assertEqual(contexto['contrato']['numero'], 'CT-0001')
        self.assertIn('data_atual', contexto['data'])

    def test_renderizar_modelo_substitui_variavel(self):
        html = renderizar_modelo('<p>{{ inquilino.nome }}</p>', self.contrato)

        self.assertEqual(html, '<p>Rodrigo Oliveira</p>')

    def test_renderizar_modelo_bloqueia_tag_de_logica(self):
        with self.assertRaises(ValueError):
            renderizar_modelo('{% if 1 %}x{% endif %}', self.contrato)


class ViewsDocumentosTests(TenantTestCase):

    def setUp(self):
        User = get_user_model()
        self.user = User.objects.create_user(username='tester', password='senha123')
        self.client.login(username='tester', password='senha123')

        self.modelo = ModeloDocumento.objects.create(
            titulo='Modelo Teste', tipo='contrato', conteudo_html='<p>{{ contrato.numero }}</p>',
        )

    def test_lista_modelos_retorna_200(self):
        resp = self.client.get(reverse('documentos:lista_modelos'), HTTP_HOST=self.domain.domain)

        self.assertEqual(resp.status_code, 200)

    def test_salvar_modelo_retorna_ok_true(self):
        resp = self.client.post(
            reverse('documentos:salvar_modelo', args=[self.modelo.pk]),
            data=json.dumps({'conteudo_html': '<p>Novo conteúdo</p>'}),
            content_type='application/json',
            HTTP_HOST=self.domain.domain,
        )

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json(), {'ok': True})
        self.modelo.refresh_from_db()
        self.assertEqual(self.modelo.conteudo_html, '<p>Novo conteúdo</p>')

    def test_editor_modelo_retorna_200_com_contexto_esperado(self):
        resp = self.client.get(
            reverse('documentos:editor_modelo', args=[self.modelo.pk]), HTTP_HOST=self.domain.domain,
        )

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context['modelo'], self.modelo)
        self.assertIn('variaveis', resp.context)
        conteudo_decodificado = json.loads(resp.context['conteudo_html_json'])
        self.assertEqual(conteudo_decodificado, self.modelo.conteudo_html)


class GerarDocumentoDownloadTests(TenantTestCase):

    def setUp(self):
        User = get_user_model()
        self.user = User.objects.create_user(username='tester', password='senha123')
        self.client.login(username='tester', password='senha123')

        self.imovel = Imovel.objects.create(
            codigo='IM-0001', tipo='apartamento', cep='60000000',
            logradouro='Rua Teste', numero='100', bairro='Centro',
            cidade='Fortaleza', estado='CE',
        )
        self.inquilino = Inquilino.objects.create(
            tipo='pf', nome='Rodrigo Oliveira', cpf='02738306006',
            telefone='85999999999', email='pagador@email.com',
            logradouro='Rua Doutor Vargas', numero='150',
            cidade='Porto Alegre', estado='RS', cep='91250000',
        )
        self.contrato = Contrato.objects.create(
            imovel=self.imovel, inquilino=self.inquilino, numero='CT-0001',
            data_inicio=date(2026, 1, 10), data_fim=date(2026, 12, 10),
            dia_vencimento=10, valor_aluguel=Decimal('1500.00'),
        )
        self.modelo = ModeloDocumento.objects.create(
            titulo='Modelo Teste', tipo='contrato', conteudo_html='<p>{{ contrato.numero }}</p>',
        )

    def test_gerar_documento_retorna_pdf(self):
        resp = self.client.post(
            reverse('documentos:gerar_documento'),
            data=json.dumps({'modelo_id': str(self.modelo.pk), 'contrato_id': self.contrato.pk}),
            content_type='application/json',
            HTTP_HOST=self.domain.domain,
        )

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertTrue(ContratoDocumentoGerado.objects.filter(contrato=self.contrato).exists())

    def test_download_documento_retorna_200_com_pdf(self):
        documento = salvar_documento_gerado(self.contrato, self.modelo, self.user)

        resp = self.client.get(
            reverse('documentos:download_documento', args=[documento.pk]), HTTP_HOST=self.domain.domain,
        )

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')


class StorageTemporarioMixin:
    """Isola o storage (B2/S3) usando FileSystemStorage em diretório temporário."""

    def setUp(self):
        super().setUp()
        pasta = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, pasta, True)
        ov = override_settings(MEDIA_ROOT=pasta)
        ov.enable()
        self.addCleanup(ov.disable)


class CriarDocumentosPadraoTests(StorageTemporarioMixin, TenantTestCase):

    TITULOS_DOCX = [d['titulo'] for d in MODELOS_PADRAO_DOCX.values()]

    def test_criar_documentos_padrao_idempotente(self):
        criar_documentos_padrao()
        variaveis1 = VariavelDocumento.objects.count()
        modelos1 = ModeloDocumento.objects.count()
        criar_documentos_padrao()

        self.assertEqual(VariavelDocumento.objects.count(), variaveis1)
        self.assertEqual(ModeloDocumento.objects.count(), modelos1)

    def test_criar_documentos_padrao_cria_modelos_docx_com_arquivo(self):
        resultado = criar_documentos_padrao()

        modelos = ModeloDocumento.objects.filter(padrao=True)
        self.assertEqual(modelos.count(), 3)
        self.assertCountEqual([m.titulo for m in modelos], self.TITULOS_DOCX)
        self.assertCountEqual([m.tipo for m in modelos], ['contrato', 'distrato', 'recibo'])
        for modelo in modelos:
            self.assertTrue(modelo.ativo)
            self.assertTrue(modelo.arquivo)
            self.assertEqual(modelo.conteudo_html, '')
            self.assertTrue(modelo.arquivo.name.endswith('.docx'))
            with modelo.arquivo.storage.open(modelo.arquivo.name, 'rb') as f:
                self.assertEqual(f.read(), caminho_modelo_padrao_docx(modelo.tipo).read_bytes())
        self.assertEqual(resultado['variaveis_criadas'], 32)
        self.assertCountEqual(resultado['modelos_criados'], self.TITULOS_DOCX)
        self.assertEqual(resultado['modelos_existentes'], [])

    def test_segunda_chamada_nao_cria_nada(self):
        criar_documentos_padrao()

        resultado = criar_documentos_padrao()

        self.assertEqual(resultado['variaveis_criadas'], 0)
        self.assertEqual(resultado['modelos_criados'], [])
        self.assertCountEqual(resultado['modelos_existentes'], self.TITULOS_DOCX)

    def test_idempotencia_por_item_recria_so_o_que_falta(self):
        criar_documentos_padrao()
        ModeloDocumento.objects.get(titulo=MODELOS_PADRAO_DOCX['distrato']['titulo']).delete()
        VariavelDocumento.objects.get(slug='inquilino.nome').delete()

        resultado = criar_documentos_padrao()

        self.assertEqual(resultado['variaveis_criadas'], 1)
        self.assertEqual(resultado['modelos_criados'], [MODELOS_PADRAO_DOCX['distrato']['titulo']])
        self.assertEqual(ModeloDocumento.objects.filter(padrao=True).count(), 3)

    def test_modelo_html_legado_padrao_nao_bloqueia_a_criacao_dos_docx(self):
        ModeloDocumento.objects.create(
            titulo='Contrato de Locação Residencial', tipo='contrato', padrao=True, conteudo_html='<p>x</p>',
        )

        criar_documentos_padrao()

        self.assertEqual(ModeloDocumento.objects.filter(padrao=True).count(), 4)

    def test_nao_altera_variavel_existente(self):
        VariavelDocumento.objects.create(
            slug='inquilino.nome', label='Rótulo do tenant', categoria='inquilino', ativo=False,
        )

        criar_documentos_padrao()

        variavel = VariavelDocumento.objects.get(slug='inquilino.nome')
        self.assertEqual(variavel.label, 'Rótulo do tenant')
        self.assertFalse(variavel.ativo)

    def test_dry_run_nao_grava(self):
        resultado = criar_modelos_padrao_docx(dry_run=True)

        self.assertCountEqual(resultado['criados'], self.TITULOS_DOCX)
        self.assertEqual(ModeloDocumento.objects.count(), 0)


class VariaveisDocumentoViewTests(StorageTemporarioMixin, TenantTestCase):

    def setUp(self):
        User = get_user_model()
        User.objects.create_user(username='tester', password='senha123')
        criar_documentos_padrao()

    def test_usuario_autorizado_ve_variaveis_agrupadas(self):
        self.client.login(username='tester', password='senha123')
        resp = self.client.get(reverse('documentos:variaveis'), HTTP_HOST=self.domain.domain)

        self.assertEqual(resp.status_code, 200)
        html = resp.content.decode()
        for variavel in VariavelDocumento.objects.filter(ativo=True):
            self.assertIn(variavel.slug, html)
        self.assertIn('Inquilino', html)
        self.assertEqual(len(resp.context['categorias']), 6)

    def test_anonimo_e_redirecionado_para_login(self):
        resp = self.client.get(reverse('documentos:variaveis'), HTTP_HOST=self.domain.domain)

        self.assertEqual(resp.status_code, 302)
        self.assertIn('login', resp['Location'])


class CatalogoXResolvedorTests(StorageTemporarioMixin, TenantTestCase):
    """Catálogo (VariavelDocumento) x chaves resolvidas por construir_contexto()."""

    def test_slugs_do_catalogo_batem_com_o_resolvedor(self):
        criar_documentos_padrao()
        proprietario = Proprietario.objects.create(nome='Maria Souza', cpf_cnpj='11122233344')
        imovel = Imovel.objects.create(
            codigo='IM-0001', tipo='apartamento', cep='60000000',
            logradouro='Rua Teste', numero='100', bairro='Centro',
            cidade='Fortaleza', estado='CE', proprietario=proprietario,
        )
        inquilino = Inquilino.objects.create(
            tipo='pf', nome='Rodrigo Oliveira', cpf='02738306006',
            telefone='85999999999', email='pagador@email.com',
            logradouro='Rua Doutor Vargas', numero='150',
            cidade='Porto Alegre', estado='RS', cep='91250000',
        )
        contrato = Contrato.objects.create(
            imovel=imovel, inquilino=inquilino, numero='CT-0001',
            data_inicio=date(2026, 1, 10), data_fim=date(2026, 12, 10),
            dia_vencimento=10, valor_aluguel=Decimal('1500.00'),
        )

        resolvidos = {
            f'{grupo}.{chave}'
            for grupo, itens in construir_contexto(contrato).items()
            for chave in itens
        }
        catalogo = set(VariavelDocumento.objects.filter(ativo=True).values_list('slug', flat=True))

        self.assertEqual(len(catalogo), 32)
        so_catalogo = sorted(catalogo - resolvidos)
        so_resolvedor = sorted(resolvidos - catalogo)
        self.assertFalse(
            so_catalogo or so_resolvedor,
            f'\nSlugs do catálogo SEM chave no resolvedor: {so_catalogo}'
            f'\nChaves do resolvedor SEM slug no catálogo: {so_resolvedor}',
        )


# ───────────────────────── Fase 2: upload de modelos .docx ─────────────────────────

def montar_docx(paragrafos=(), header=None, footer=None):
    """Gera um .docx em memória. Cada parágrafo é uma lista de runs (texto fragmentado)."""
    doc = Document()
    for runs in paragrafos:
        p = doc.add_paragraph()
        for texto in runs:
            p.add_run(texto)
    if header:
        doc.sections[0].header.paragraphs[0].text = header
    if footer:
        doc.sections[0].footer.paragraphs[0].text = footer
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def montar_zip(entradas):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        for nome, conteudo in entradas.items():
            zf.writestr(nome, conteudo)
    return buf.getvalue()


def upload(conteudo, nome='modelo.docx'):
    return SimpleUploadedFile(nome, conteudo)


class ValidatorsDocxTests(SimpleTestCase):

    def test_docx_valido_passa(self):
        ValidarDocx()(upload(montar_docx([['Olá']])))

    def test_nao_zip_falha(self):
        with self.assertRaises(ValidationError):
            ValidarDocx()(upload(b'isto nao e um zip'))

    def test_pk_sem_zip_valido_falha(self):
        with self.assertRaises(ValidationError):
            ValidarDocx()(upload(b'PK\x03\x04lixo'))

    def test_zip_sem_partes_do_word_falha(self):
        with self.assertRaises(ValidationError):
            ValidarDocx()(upload(montar_zip({'qualquer.txt': 'x'})))

    def test_extensao_errada_falha(self):
        with self.assertRaises(ValidationError):
            validar_extensao_docx(upload(montar_docx(), nome='modelo.pdf'))

    def test_acima_do_limite_de_tamanho_falha(self):
        with self.assertRaisesMessage(ValidationError, 'Arquivo acima de 2 MB.'):
            validar_tamanho_modelo(upload(b'0' * (2 * 1024 * 1024 + 1)))

    def test_abaixo_do_limite_de_tamanho_passa(self):
        validar_tamanho_modelo(upload(b'0' * 1024))

    @override_settings(DOCUMENTO_MODELO_MAX_DESCOMPACTADO_MB=1)
    def test_zip_bomb_falha(self):
        conteudo = montar_zip({
            '[Content_Types].xml': '<x/>',
            'word/document.xml': '<x/>',
            'word/enorme.bin': b'\0' * (3 * 1024 * 1024),
        })
        with self.assertRaises(ValidationError):
            ValidarDocx()(upload(conteudo))

    @override_settings(DOCUMENTO_MODELO_MAX_ENTRADAS_ZIP=3)
    def test_entradas_demais_falha(self):
        conteudo = montar_zip({
            '[Content_Types].xml': '<x/>', 'word/document.xml': '<x/>',
            'a.xml': '<x/>', 'b.xml': '<x/>',
        })
        with self.assertRaises(ValidationError):
            ValidarDocx()(upload(conteudo))

    def test_macro_vba_falha(self):
        conteudo = montar_zip({
            '[Content_Types].xml': '<x/>', 'word/document.xml': '<x/>',
            'word/vbaProject.bin': b'macro',
        })
        with self.assertRaisesMessage(ValidationError, 'macros'):
            ValidarDocx()(upload(conteudo))

    def test_entrada_com_path_traversal_falha(self):
        conteudo = montar_zip({
            '[Content_Types].xml': '<x/>', 'word/document.xml': '<x/>',
            '../evil.txt': 'x',
        })
        with self.assertRaises(ValidationError):
            ValidarDocx()(upload(conteudo))


class AnalisarDocxTests(StorageTemporarioMixin, TenantTestCase):

    def setUp(self):
        criar_documentos_padrao()

    def test_variavel_integra(self):
        r = analisar_docx(io.BytesIO(montar_docx([['Nome: {{ inquilino.nome }}']])))

        self.assertEqual(r['variaveis'], ['inquilino.nome'])
        self.assertEqual(r['desconhecidas'], [])
        self.assertEqual(r['erros'], [])

    def test_variavel_fragmentada_em_varios_runs(self):
        r = analisar_docx(io.BytesIO(montar_docx([['{{ inqui', 'lino.n', 'ome }}']])))

        self.assertEqual(r['variaveis'], ['inquilino.nome'])
        self.assertEqual(r['erros'], [])

    def test_variavel_em_header_e_footer(self):
        r = analisar_docx(io.BytesIO(montar_docx(
            [['corpo']], header='{{ contrato.numero }}', footer='{{ data.data_atual }}',
        )))

        self.assertEqual(r['variaveis'], ['contrato.numero', 'data.data_atual'])

    def test_variavel_desconhecida(self):
        r = analisar_docx(io.BytesIO(montar_docx([['{{ inquilino.nome }} {{ foo.bar }}']])))

        self.assertEqual(r['desconhecidas'], ['foo.bar'])
        self.assertEqual(r['erros'], [])

    def test_tag_de_logica_gera_erro(self):
        r = analisar_docx(io.BytesIO(montar_docx([['{% if x %}oi{% endif %}']])))

        self.assertTrue(r['erros'])

    def test_abre_chave_sem_fechar_gera_erro(self):
        r = analisar_docx(io.BytesIO(montar_docx([['Nome {{ inquilino.nome']])))

        self.assertTrue(r['erros'])


class UploadModeloViewsTests(StorageTemporarioMixin, TenantTestCase):

    def setUp(self):
        super().setUp()
        criar_documentos_padrao()
        User = get_user_model()
        User.objects.create_user(username='tester', password='senha123')
        self.client.login(username='tester', password='senha123')
        self.host = self.domain.domain

    def _post_criar(self, conteudo, titulo='Contrato Novo', nome='modelo.docx'):
        return self.client.post(
            reverse('documentos:criar_modelo'),
            data={'titulo': titulo, 'tipo': 'contrato', 'arquivo': upload(conteudo, nome)},
            HTTP_HOST=self.host,
        )

    def _mensagens(self, resp):
        return [str(m) for m in resp.context['messages']]

    def test_criar_com_upload_valido_salva_e_redireciona_para_lista(self):
        resp = self._post_criar(montar_docx([['{{ inquilino.nome }}']]))

        self.assertRedirects(resp, reverse('documentos:lista_modelos'), fetch_redirect_response=False)
        modelo = ModeloDocumento.objects.get(titulo='Contrato Novo')
        self.assertTrue(modelo.arquivo)
        self.assertTrue(modelo.arquivo.name.endswith('.docx'))
        self.assertIn('/documentos/modelos/', modelo.arquivo.name)
        self.assertTrue(modelo.arquivo.storage.exists(modelo.arquivo.name))

    def test_variavel_desconhecida_gera_aviso(self):
        resp = self.client.post(
            reverse('documentos:criar_modelo'),
            data={'titulo': 'Com Aviso', 'tipo': 'outro',
                  'arquivo': upload(montar_docx([['{{ foo.bar }}']]))},
            HTTP_HOST=self.host, follow=True,
        )

        self.assertTrue(ModeloDocumento.objects.filter(titulo='Com Aviso').exists())
        self.assertTrue(any('foo.bar' in m for m in self._mensagens(resp)))

    def test_upload_com_tag_de_logica_nao_cria_registro(self):
        resp = self._post_criar(montar_docx([['{% if x %}oi{% endif %}']]), titulo='Invalido')

        self.assertEqual(resp.status_code, 200)
        self.assertFalse(ModeloDocumento.objects.filter(titulo='Invalido').exists())
        self.assertTrue(resp.context['form_novo'].errors['arquivo'])

    def test_upload_nao_docx_nao_cria_registro(self):
        resp = self._post_criar(b'nao sou docx', titulo='Lixo')

        self.assertEqual(resp.status_code, 200)
        self.assertFalse(ModeloDocumento.objects.filter(titulo='Lixo').exists())

    def test_criar_sem_arquivo_nao_cria_registro(self):
        resp = self.client.post(
            reverse('documentos:criar_modelo'),
            data={'titulo': 'Sem Arquivo', 'tipo': 'outro'}, HTTP_HOST=self.host,
        )

        self.assertEqual(resp.status_code, 200)
        self.assertFalse(ModeloDocumento.objects.filter(titulo='Sem Arquivo').exists())

    def test_download_retorna_o_arquivo(self):
        conteudo = montar_docx([['{{ inquilino.nome }}']])
        self._post_criar(conteudo, titulo='Meu Contrato')
        modelo = ModeloDocumento.objects.get(titulo='Meu Contrato')

        resp = self.client.get(reverse('documentos:download_modelo', args=[modelo.pk]), HTTP_HOST=self.host)

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(b''.join(resp.streaming_content), conteudo)
        self.assertIn('meu-contrato.docx', resp['Content-Disposition'])

    def test_download_de_modelo_sem_arquivo_da_404(self):
        modelo = ModeloDocumento.objects.create(titulo='Legado', tipo='outro')

        resp = self.client.get(reverse('documentos:download_modelo', args=[modelo.pk]), HTTP_HOST=self.host)

        self.assertEqual(resp.status_code, 404)

    def test_substituir_arquivo_troca_e_remove_o_antigo(self):
        self._post_criar(montar_docx([['antigo']]), titulo='Trocar')
        modelo = ModeloDocumento.objects.get(titulo='Trocar')
        nome_antigo = modelo.arquivo.name
        novo = montar_docx([['{{ contrato.numero }}']])

        resp = self.client.post(
            reverse('documentos:substituir_arquivo_modelo', args=[modelo.pk]),
            data={'arquivo': upload(novo)}, HTTP_HOST=self.host,
        )

        self.assertRedirects(resp, reverse('documentos:lista_modelos'), fetch_redirect_response=False)
        modelo.refresh_from_db()
        self.assertNotEqual(modelo.arquivo.name, nome_antigo)
        self.assertTrue(modelo.arquivo.storage.exists(modelo.arquivo.name))
        self.assertFalse(modelo.arquivo.storage.exists(nome_antigo))

    def test_substituir_com_arquivo_invalido_mantem_o_antigo(self):
        self._post_criar(montar_docx([['antigo']]), titulo='Manter')
        modelo = ModeloDocumento.objects.get(titulo='Manter')
        nome_antigo = modelo.arquivo.name

        self.client.post(
            reverse('documentos:substituir_arquivo_modelo', args=[modelo.pk]),
            data={'arquivo': upload(b'lixo')}, HTTP_HOST=self.host,
        )

        modelo.refresh_from_db()
        self.assertEqual(modelo.arquivo.name, nome_antigo)
        self.assertTrue(modelo.arquivo.storage.exists(nome_antigo))

    def test_anonimo_e_redirecionado(self):
        modelo = ModeloDocumento.objects.create(titulo='X', tipo='outro')
        self.client.logout()

        for nome, args in [('criar_modelo', []), ('download_modelo', [modelo.pk]),
                           ('substituir_arquivo_modelo', [modelo.pk])]:
            resp = self.client.post(reverse(f'documentos:{nome}', args=args), HTTP_HOST=self.host)
            self.assertEqual(resp.status_code, 302, nome)
            self.assertIn('login', resp['Location'], nome)


# ───────────────────────── Fase 3A: geração de PDF a partir do .docx ─────────────────────────

def criar_contrato_teste():
    proprietario = Proprietario.objects.create(nome='Maria Souza', cpf_cnpj='11122233344')
    imovel = Imovel.objects.create(
        codigo='IM-0001', tipo='apartamento', cep='60000000',
        logradouro='Rua Teste', numero='100', bairro='Centro',
        cidade='Fortaleza', estado='CE', proprietario=proprietario,
    )
    inquilino = Inquilino.objects.create(
        tipo='pf', nome='Rodrigo Oliveira', cpf='02738306006',
        telefone='85999999999', email='pagador@email.com',
        logradouro='Rua Doutor Vargas', numero='150',
        cidade='Porto Alegre', estado='RS', cep='91250000',
    )
    return Contrato.objects.create(
        imovel=imovel, inquilino=inquilino, numero='CT-0001',
        data_inicio=date(2026, 1, 10), data_fim=date(2026, 12, 10),
        dia_vencimento=10, valor_aluguel=Decimal('1500.00'),
    )


def criar_modelo_docx(conteudo, titulo='Modelo Docx'):
    modelo = ModeloDocumento(titulo=titulo, tipo='contrato')
    modelo.arquivo.save('modelo.docx', ContentFile(conteudo), save=False)
    modelo.save()
    return modelo


def textos_docx(docx_bytes):
    """Texto do corpo, cabeçalho e rodapé de um .docx renderizado."""
    doc = Document(io.BytesIO(docx_bytes))
    sec = doc.sections[0]
    return {
        'corpo': '\n'.join(p.text for p in doc.paragraphs),
        'header': sec.header.paragraphs[0].text,
        'footer': sec.footer.paragraphs[0].text,
    }


class FmtMoneyTests(SimpleTestCase):

    def test_formata_no_padrao_brasileiro(self):
        self.assertEqual(_fmt_money(Decimal('1234.5')), 'R$ 1.234,50')

    def test_zero(self):
        self.assertEqual(_fmt_money(Decimal('0')), 'R$ 0,00')

    def test_milhoes(self):
        self.assertEqual(_fmt_money(Decimal('1000000')), 'R$ 1.000.000,00')

    def test_nulo_retorna_vazio(self):
        self.assertEqual(_fmt_money(None), '')

    def test_arredonda_centavos(self):
        self.assertEqual(_fmt_money(Decimal('10.005')), 'R$ 10,01')


class RenderizarDocxTests(StorageTemporarioMixin, TenantTestCase):

    CTX = {'inquilino': {'nome': 'João & <b>Silva</b>'}, 'contrato': {'numero': 'CT-9'}}

    def _render(self, conteudo, ctx=None):
        return renderizar_docx(criar_modelo_docx(conteudo), ctx or self.CTX)

    def test_variavel_integra(self):
        out = self._render(montar_docx([['Contrato {{ contrato.numero }}']]))

        self.assertIn('Contrato CT-9', textos_docx(out)['corpo'])

    def test_variavel_fragmentada_em_varios_runs(self):
        out = self._render(montar_docx([['{{ contra', 'to.nu', 'mero }}']]))

        self.assertIn('CT-9', textos_docx(out)['corpo'])
        self.assertNotIn('{{', textos_docx(out)['corpo'])

    def test_variavel_em_header_e_footer(self):
        out = self._render(montar_docx(
            [['corpo']], header='H {{ contrato.numero }}', footer='F {{ contrato.numero }}',
        ))
        t = textos_docx(out)

        self.assertEqual(t['header'], 'H CT-9')
        self.assertEqual(t['footer'], 'F CT-9')

    def test_variavel_ausente_renderiza_vazio(self):
        out = self._render(montar_docx([['[{{ contrato.nao_existe }}]']]))

        self.assertIn('[]', textos_docx(out)['corpo'])

    def test_autoescape_de_e_comercial_e_menor_que(self):
        out = self._render(montar_docx([['{{ inquilino.nome }}']]))

        self.assertEqual(textos_docx(out)['corpo'].strip().splitlines()[-1], 'João & <b>Silva</b>')

    def test_sandbox_bloqueia_acesso_a_atributo_perigoso(self):
        conteudo = montar_docx([["{{ ''.__class__.__mro__ }}"]])

        with self.assertRaises(SecurityError):
            self._render(conteudo)

    def test_tag_de_logica_e_rejeitada(self):
        with self.assertRaises(ModeloDocxInvalido):
            self._render(montar_docx([['{% if x %}oi{% endif %}']]))

    def test_tag_de_logica_fragmentada_em_runs_e_rejeitada(self):
        with self.assertRaises(ModeloDocxInvalido):
            self._render(montar_docx([['{', '% if x %', '}oi']]))


class ConverterDocxParaPdfTests(SimpleTestCase):

    @patch('apps.documentos.services.os.killpg', create=True)
    @patch('apps.documentos.services.subprocess.Popen')
    def test_timeout_mata_o_grupo_e_levanta_erro(self, popen, killpg):
        proc = popen.return_value
        proc.pid = 4242
        proc.communicate.side_effect = [subprocess.TimeoutExpired('soffice', 60), (b'', b'')]

        with self.assertRaises(ConversaoDocxErro):
            converter_docx_para_pdf(b'x')

        self.assertTrue(popen.call_args.kwargs['start_new_session'])
        self.assertIsInstance(popen.call_args.args[0], list)
        killpg.assert_called_once()

    @patch('apps.documentos.services.subprocess.Popen')
    def test_falha_do_soffice_levanta_erro_sem_vazar_stderr(self, popen):
        proc = popen.return_value
        proc.returncode = 1
        proc.communicate.return_value = (b'', b'SEGREDO-DO-STDERR')

        with self.assertRaises(ConversaoDocxErro) as ctx:
            converter_docx_para_pdf(b'x')

        self.assertNotIn('SEGREDO', str(ctx.exception))

    @patch('apps.documentos.services.subprocess.Popen', side_effect=FileNotFoundError)
    def test_soffice_ausente_levanta_erro(self, popen):
        with self.assertRaises(ConversaoDocxErro):
            converter_docx_para_pdf(b'x')

    @patch('apps.documentos.services.shutil.rmtree')
    @patch('apps.documentos.services.subprocess.Popen')
    def test_limpa_tmp_mesmo_com_erro(self, popen, rmtree):
        popen.return_value.returncode = 1
        popen.return_value.communicate.return_value = (b'', b'')

        with self.assertRaises(ConversaoDocxErro):
            converter_docx_para_pdf(b'x')

        rmtree.assert_called_once()

    @skipUnless(shutil.which('soffice'), 'soffice não instalado')
    def test_integracao_com_soffice_real(self):
        pdf = converter_docx_para_pdf(montar_docx([['Olá mundo']]))

        self.assertTrue(pdf.startswith(b'%PDF'))


class TaskGerarDocumentoDocxTests(StorageTemporarioMixin, TenantTestCase):

    def setUp(self):
        super().setUp()
        self.contrato = criar_contrato_teste()
        sentry = patch('apps.documentos.tasks.sentry_sdk')
        self.sentry = sentry.start()
        self.addCleanup(sentry.stop)

    def _documento(self, conteudo):
        modelo = criar_modelo_docx(conteudo)
        return ContratoDocumentoGerado.objects.create(
            contrato=self.contrato, modelo=modelo, titulo='Doc', status='pendente',
        )

    def _rodar(self, documento):
        gerar_documento_docx(self.tenant.schema_name, str(documento.pk))
        documento.refresh_from_db()
        return documento

    @patch('apps.documentos.tasks.converter_docx_para_pdf', return_value=b'%PDF-fake')
    def test_sucesso(self, conv):
        doc = self._rodar(self._documento(montar_docx([['{{ contrato.numero }}']])))

        self.assertEqual(doc.status, 'gerado')
        self.assertEqual(doc.erro_msg, '')
        self.assertEqual(doc.arquivo_pdf.read(), b'%PDF-fake')
        docx_enviado = conv.call_args.args[0]
        self.assertIn('CT-0001', textos_docx(docx_enviado)['corpo'])

    @patch('apps.documentos.tasks.converter_docx_para_pdf')
    def test_falha_na_renderizacao(self, conv):
        doc = self._rodar(self._documento(montar_docx([['{% if x %}oi{% endif %}']])))

        self.assertEqual(doc.status, 'erro')
        self.assertIn('lógica', doc.erro_msg)
        conv.assert_not_called()
        self.sentry.capture_exception.assert_called_once()

    @patch('apps.documentos.tasks.converter_docx_para_pdf', side_effect=ConversaoDocxErro('Falha ao converter o documento para PDF.'))
    def test_erro_de_conversao(self, conv):
        doc = self._rodar(self._documento(montar_docx([['oi']])))

        self.assertEqual(doc.status, 'erro')
        self.assertEqual(doc.erro_msg, 'Falha ao converter o documento para PDF.')

    @patch('apps.documentos.tasks.converter_docx_para_pdf', side_effect=SoftTimeLimitExceeded())
    def test_timeout_da_task(self, conv):
        doc = self._rodar(self._documento(montar_docx([['oi']])))

        self.assertEqual(doc.status, 'erro')
        self.assertEqual(doc.erro_msg, ERRO_GENERICO)

    @patch('apps.documentos.tasks.converter_docx_para_pdf', side_effect=RuntimeError('detalhe interno sensível'))
    def test_erro_inesperado_nao_vaza_detalhe(self, conv):
        doc = self._rodar(self._documento(montar_docx([['oi']])))

        self.assertEqual(doc.status, 'erro')
        self.assertNotIn('sensível', doc.erro_msg)

    def test_documento_inexistente_nao_levanta(self):
        gerar_documento_docx(self.tenant.schema_name, '00000000-0000-0000-0000-000000000000')

    def test_task_roteada_para_fila_docx(self):
        from django.conf import settings
        self.assertEqual(settings.CELERY_TASK_ROUTES['apps.documentos.tasks.*'], {'queue': 'docx'})


class ViewsGeracaoDocxTests(StorageTemporarioMixin, TenantTestCase):

    def setUp(self):
        super().setUp()
        criar_documentos_padrao()
        User = get_user_model()
        self.user = User.objects.create_user(username='tester', password='senha123')
        self.client.login(username='tester', password='senha123')
        self.host = self.domain.domain
        self.contrato = criar_contrato_teste()
        self.modelo_docx = criar_modelo_docx(montar_docx([['{{ contrato.numero }}']]))

    def _gerar(self, modelo):
        return self.client.post(
            reverse('documentos:gerar_documento'),
            data=json.dumps({'modelo_id': str(modelo.pk), 'contrato_id': self.contrato.pk}),
            content_type='application/json', HTTP_HOST=self.host,
        )

    def _doc(self, status, **kw):
        return ContratoDocumentoGerado.objects.create(
            contrato=self.contrato, modelo=self.modelo_docx, titulo='Doc', status=status, **kw,
        )

    @patch('apps.documentos.views.gerar_documento_docx.delay')
    def test_modelo_com_arquivo_retorna_202_e_dispara_task(self, delay):
        resp = self._gerar(self.modelo_docx)

        self.assertEqual(resp.status_code, 202)
        doc = ContratoDocumentoGerado.objects.get()
        self.assertEqual(doc.status, 'pendente')
        self.assertEqual(doc.gerado_por, self.user)
        self.assertEqual(resp.json()['id'], str(doc.pk))
        self.assertEqual(resp.json()['status_url'], reverse('documentos:status_documento', args=[doc.pk]))
        delay.assert_called_once_with(self.tenant.schema_name, str(doc.pk))

    @patch('apps.documentos.views.gerar_documento_docx.delay')
    def test_fluxo_legado_continua_sincrono(self, delay):
        legado = ModeloDocumento.objects.create(
            titulo='Legado', tipo='contrato', conteudo_html='<p>{{ contrato.numero }}</p>',
        )

        resp = self._gerar(legado)

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        delay.assert_not_called()
        self.assertEqual(ContratoDocumentoGerado.objects.get().status, 'gerado')

    def test_status_gerado_traz_download_url(self):
        doc = self._doc('gerado')
        doc.arquivo_pdf.save('documento.pdf', ContentFile(b'%PDF'), save=True)

        resp = self.client.get(reverse('documentos:status_documento', args=[doc.pk]), HTTP_HOST=self.host)

        self.assertEqual(resp.json()['status'], 'gerado')
        self.assertEqual(resp.json()['download_url'], reverse('documentos:download_documento', args=[doc.pk]))

    def test_status_pendente_recente_continua_pendente(self):
        doc = self._doc('pendente')

        resp = self.client.get(reverse('documentos:status_documento', args=[doc.pk]), HTTP_HOST=self.host)

        self.assertEqual(resp.json()['status'], 'pendente')
        self.assertIsNone(resp.json()['download_url'])

    def test_status_preso_ha_mais_de_5_minutos_vira_erro(self):
        for status in ('pendente', 'processando'):
            doc = self._doc(status)
            ContratoDocumentoGerado.objects.filter(pk=doc.pk).update(
                gerado_em=timezone.now() - timedelta(minutes=6),
            )

            resp = self.client.get(reverse('documentos:status_documento', args=[doc.pk]), HTTP_HOST=self.host)

            self.assertEqual(resp.json()['status'], 'erro', status)
            self.assertIn('Tempo esgotado', resp.json()['erro'])
            doc.refresh_from_db()
            self.assertEqual(doc.status, 'erro')

    def test_download_so_quando_gerado(self):
        doc = self._doc('pendente')
        doc.arquivo_pdf.save('documento.pdf', ContentFile(b'%PDF'), save=True)

        resp = self.client.get(reverse('documentos:download_documento', args=[doc.pk]), HTTP_HOST=self.host)
        self.assertEqual(resp.status_code, 404)

        doc.status = 'gerado'
        doc.save()
        resp = self.client.get(reverse('documentos:download_documento', args=[doc.pk]), HTTP_HOST=self.host)
        self.assertEqual(resp.status_code, 200)

    def test_status_anonimo_redireciona(self):
        doc = self._doc('pendente')
        self.client.logout()

        resp = self.client.get(reverse('documentos:status_documento', args=[doc.pk]), HTTP_HOST=self.host)

        self.assertEqual(resp.status_code, 302)


class ModelosPadraoDocxPacoteTests(StorageTemporarioMixin, TenantTestCase):
    """Os .docx versionados em apps/documentos/modelos_padrao/."""

    def setUp(self):
        super().setUp()
        criar_variaveis_padrao()
        proprietario = Proprietario.objects.create(nome='Maria Souza', cpf_cnpj='11122233344')
        imovel = Imovel.objects.create(
            codigo='IM-0001', tipo='apartamento', cep='60000000',
            logradouro='Rua Teste', numero='100', bairro='Centro',
            cidade='Fortaleza', estado='CE', proprietario=proprietario,
        )
        inquilino = Inquilino.objects.create(
            tipo='pf', nome='Rodrigo Oliveira', cpf='02738306006',
            telefone='85999999999', email='pagador@email.com',
            logradouro='Rua Doutor Vargas', numero='150',
            cidade='Porto Alegre', estado='RS', cep='91250000',
        )
        self.contrato = Contrato.objects.create(
            imovel=imovel, inquilino=inquilino, numero='CT-0001',
            data_inicio=date(2026, 1, 10), data_fim=date(2026, 12, 10),
            dia_vencimento=10, valor_aluguel=Decimal('1500.00'),
        )

    def _bytes(self, tipo):
        return caminho_modelo_padrao_docx(tipo).read_bytes()

    def test_pacote_tem_os_tres_arquivos(self):
        self.assertEqual(set(MODELOS_PADRAO_DOCX), {'contrato', 'distrato', 'recibo'})
        for tipo in MODELOS_PADRAO_DOCX:
            self.assertTrue(caminho_modelo_padrao_docx(tipo).is_file(), tipo)

    def test_passa_no_validador_e_na_analise_sem_erros_nem_desconhecidas(self):
        for tipo in MODELOS_PADRAO_DOCX:
            with self.subTest(tipo=tipo):
                ValidarDocx()(upload(self._bytes(tipo)))
                analise = analisar_docx(io.BytesIO(self._bytes(tipo)))
                self.assertEqual(analise['erros'], [])
                self.assertEqual(analise['desconhecidas'], [])
                self.assertTrue(analise['variaveis'])

    def test_variavel_inteira_em_um_unico_run(self):
        for tipo in MODELOS_PADRAO_DOCX:
            with self.subTest(tipo=tipo):
                doc = Document(io.BytesIO(self._bytes(tipo)))
                for par in doc.paragraphs:
                    for run in par.runs:
                        self.assertEqual(run.text.count('{{'), run.text.count('}}'), run.text)
                        if '{{' in run.text:
                            self.assertRegex(run.text, r'^\{\{ [\w.]+ \}\}$')

    def test_texto_identico_ao_da_fixture(self):
        fixture = json.loads(
            (caminho_modelo_padrao_docx('contrato').parent.parent / 'fixtures' / 'modelos_padrao.json')
            .read_text(encoding='utf-8')
        )
        for item in fixture:
            tipo = item['fields']['tipo']
            with self.subTest(tipo=tipo):
                esperado = [
                    html.unescape(re.sub(r'<[^>]+>', '', trecho.replace('<br>', '\n')))
                    for trecho in re.findall(r'<(?:h1|p)>(.*?)</(?:h1|p)>', item['fields']['conteudo_html'])
                ]
                doc = Document(io.BytesIO(self._bytes(tipo)))
                self.assertEqual([p.text for p in doc.paragraphs], esperado)

    def test_renderizar_com_contexto_completo_nao_deixa_chaves(self):
        contexto = construir_contexto(self.contrato)
        chaves = {f'{grupo}.{chave}' for grupo, itens in contexto.items() for chave in itens}
        self.assertEqual(len(chaves), 32)
        criar_modelos_padrao_docx()

        for modelo in ModeloDocumento.objects.filter(padrao=True):
            with self.subTest(modelo=modelo.titulo):
                saida = renderizar_docx(modelo, contexto)
                doc = Document(io.BytesIO(saida))
                texto = '\n'.join(p.text for p in doc.paragraphs)
                self.assertNotIn('{{', texto)
                self.assertNotIn('}}', texto)
                self.assertIn('Rodrigo Oliveira', texto)

    @skipUnless(shutil.which(settings.DOCUMENTO_SOFFICE_BIN), 'soffice não instalado')
    def test_conversao_real_gera_pdf_valido_por_modelo(self):
        contexto = construir_contexto(self.contrato)
        criar_modelos_padrao_docx()

        for modelo in ModeloDocumento.objects.filter(padrao=True):
            with self.subTest(modelo=modelo.titulo):
                pdf = converter_docx_para_pdf(renderizar_docx(modelo, contexto))
                self.assertTrue(pdf.startswith(b'%PDF'))
                self.assertGreater(len(pdf), 1000)


class BackfillModelosPadraoDocxTests(StorageTemporarioMixin, TenantTestCase):

    def setUp(self):
        super().setUp()
        self.legados = [
            ModeloDocumento.objects.create(
                titulo=dados['titulo'].removesuffix(' (DOCX)'), tipo=tipo, padrao=True,
                conteudo_html=f'<p>legado {tipo}</p>',
            )
            for tipo, dados in MODELOS_PADRAO_DOCX.items()
        ]
        self.schema = self.tenant.schema_name

    def _rodar(self, **kwargs):
        saida = io.StringIO()
        call_command('backfill_modelos_padrao_docx', stdout=saida, stderr=io.StringIO(), **kwargs)
        return saida.getvalue()

    def _tenant_sem_schema(self, nome='imob_semschema'):
        tenant = Tenant(
            schema_name=nome, nome='Sem schema', email='x@y.com',
            plano=Plano.objects.filter(ativo=True).first(), provisionamento_status='pronto',
        )
        tenant.auto_create_schema = False
        with schema_context(get_public_schema_name()):
            tenant.save()
        return tenant

    def _snapshot_legados(self):
        return [
            (m.pk, m.titulo, m.conteudo_html, m.padrao, bool(m.arquivo), m.atualizado_em)
            for m in ModeloDocumento.objects.filter(pk__in=[l.pk for l in self.legados]).order_by('pk')
        ]

    def test_sem_schema_nem_all_aborta(self):
        with self.assertRaises(CommandError):
            self._rodar()

    def test_schema_e_all_juntos_abortam(self):
        with self.assertRaises(CommandError):
            self._rodar(schema=self.schema, all=True)

    def test_schema_inexistente_aborta(self):
        with self.assertRaises(CommandError):
            self._rodar(schema='nao_existe')

    def test_adiciona_os_tres_docx_e_preserva_os_html(self):
        antes = self._snapshot_legados()

        saida = self._rodar(schema=self.schema)

        self.assertEqual(ModeloDocumento.objects.count(), 6)
        docx = ModeloDocumento.objects.exclude(arquivo='')
        self.assertCountEqual([m.titulo for m in docx], [d['titulo'] for d in MODELOS_PADRAO_DOCX.values()])
        self.assertTrue(all(m.padrao and m.ativo for m in docx))
        self.assertEqual(self._snapshot_legados(), antes)
        self.assertIn('criados=3 já existentes=0 variáveis novas=32 erros=0', saida)

    def test_segunda_execucao_e_noop(self):
        self._rodar(schema=self.schema)
        arquivos = set(ModeloDocumento.objects.exclude(arquivo='').values_list('arquivo', flat=True))

        saida = self._rodar(schema=self.schema)

        self.assertEqual(ModeloDocumento.objects.count(), 6)
        self.assertEqual(
            set(ModeloDocumento.objects.exclude(arquivo='').values_list('arquivo', flat=True)), arquivos,
        )
        self.assertIn('criados=0 já existentes=3 variáveis novas=0 erros=0', saida)

    def test_dry_run_nao_grava(self):
        saida = self._rodar(schema=self.schema, dry_run=True)

        self.assertEqual(ModeloDocumento.objects.count(), 3)
        self.assertIn('criados=3', saida)
        self.assertIn('DRY-RUN', saida)

    def test_schema_restringe_aos_tenants_pedidos(self):
        self._tenant_sem_schema()

        self._rodar(schema=self.schema)

        self.assertEqual(ModeloDocumento.objects.count(), 6)

    def test_falha_em_um_tenant_nao_interrompe_os_outros_e_sai_com_erro(self):
        self._tenant_sem_schema('imob_a_falha')

        with self.assertRaises(CommandError) as ctx:
            self._rodar(all=True)

        self.assertIn('1 tenant(s) com erro', str(ctx.exception))
        self.assertEqual(ModeloDocumento.objects.exclude(arquivo='').count(), 3)

    def test_resumo_final_com_contagem(self):
        saida = io.StringIO()
        call_command('backfill_modelos_padrao_docx', all=True, stdout=saida, stderr=io.StringIO())

        self.assertIn('Resumo: criados=3 já existentes=0 variáveis novas=32 erros=0', saida.getvalue())


class ArquivarModelosLegadosTests(StorageTemporarioMixin, TenantTestCase):

    def setUp(self):
        super().setUp()
        self.schema = self.tenant.schema_name
        self.legados = {
            tipo: ModeloDocumento.objects.create(
                titulo=dados['titulo'].removesuffix(' (DOCX)'), tipo=tipo, padrao=True,
                conteudo_html=f'<p>legado {tipo}</p>',
            )
            for tipo, dados in MODELOS_PADRAO_DOCX.items()
        }
        self.personalizado = ModeloDocumento.objects.create(
            titulo='Meu contrato', tipo='contrato', padrao=False, conteudo_html='<p>do cliente</p>',
        )
        criar_modelos_padrao_docx()

    def _rodar(self, **kwargs):
        saida = io.StringIO()
        call_command('arquivar_modelos_legados', stdout=saida, stderr=io.StringIO(), **kwargs)
        return saida.getvalue()

    def _ativo(self, modelo):
        return ModeloDocumento.objects.get(pk=modelo.pk).ativo

    def _tenant_sem_schema(self, nome='imob_a_falha'):
        tenant = Tenant(
            schema_name=nome, nome='Sem schema', email='x@y.com',
            plano=Plano.objects.filter(ativo=True).first(), provisionamento_status='pronto',
        )
        tenant.auto_create_schema = False
        with schema_context(get_public_schema_name()):
            tenant.save()
        return tenant

    def test_sem_schema_nem_all_aborta(self):
        with self.assertRaises(CommandError):
            self._rodar()

    def test_padrao_com_equivalente_docx_e_desativado(self):
        saida = self._rodar(schema=self.schema)

        for legado in self.legados.values():
            self.assertFalse(self._ativo(legado))
        self.assertTrue(self._ativo(self.personalizado))
        self.assertEqual(ModeloDocumento.objects.filter(ativo=True).exclude(arquivo='').count(), 3)
        self.assertEqual(ModeloDocumento.objects.count(), 7)
        self.assertIn('desativados=3', saida)

    def test_sem_equivalente_docx_e_ignorado(self):
        ModeloDocumento.objects.get(titulo=MODELOS_PADRAO_DOCX['distrato']['titulo']).delete()

        saida = self._rodar(schema=self.schema)

        self.assertTrue(self._ativo(self.legados['distrato']))
        self.assertFalse(self._ativo(self.legados['contrato']))
        self.assertIn('ignorados (sem equivalente DOCX)=1', saida)
        self.assertIn('Resumo: desativados=2 ignorados=1', saida)

    def test_equivalente_docx_inativo_nao_conta(self):
        ModeloDocumento.objects.filter(titulo=MODELOS_PADRAO_DOCX['recibo']['titulo']).update(ativo=False)

        self._rodar(schema=self.schema)

        self.assertTrue(self._ativo(self.legados['recibo']))

    def test_personalizado_so_e_listado(self):
        saida = self._rodar(all=True)

        self.assertTrue(self._ativo(self.personalizado))
        self.assertIn(str(self.personalizado.pk), saida)
        self.assertIn('Meu contrato', saida)
        self.assertIn(f'conteudo_html={len(self.personalizado.conteudo_html)} chars', saida)
        self.assertIn('personalizados listados=1', saida)

    def test_incluir_personalizados_com_schema_desativa(self):
        self._rodar(schema=self.schema, incluir_personalizados=True)

        self.assertFalse(self._ativo(self.personalizado))
        self.assertEqual(ModeloDocumento.objects.count(), 7)

    def test_incluir_personalizados_com_all_aborta(self):
        with self.assertRaises(CommandError):
            self._rodar(all=True, incluir_personalizados=True)

        self.assertTrue(self._ativo(self.personalizado))
        self.assertTrue(all(self._ativo(m) for m in self.legados.values()))

    def test_segunda_execucao_e_noop(self):
        self._rodar(schema=self.schema, incluir_personalizados=True)
        estado = list(ModeloDocumento.objects.order_by('pk').values_list('pk', 'ativo', 'atualizado_em'))

        saida = self._rodar(schema=self.schema, incluir_personalizados=True)

        self.assertEqual(list(ModeloDocumento.objects.order_by('pk').values_list('pk', 'ativo', 'atualizado_em')), estado)
        self.assertIn('Resumo: desativados=0 ignorados=0 personalizados=0 erros=0', saida)

    def test_dry_run_nao_grava(self):
        estado = list(ModeloDocumento.objects.order_by('pk').values_list('pk', 'ativo', 'atualizado_em'))

        saida = self._rodar(schema=self.schema, incluir_personalizados=True, dry_run=True)

        self.assertEqual(list(ModeloDocumento.objects.order_by('pk').values_list('pk', 'ativo', 'atualizado_em')), estado)
        self.assertIn('DRY-RUN', saida)
        self.assertIn('desativados=4', saida)

    def test_falha_em_um_tenant_nao_interrompe_os_outros_e_sai_com_erro(self):
        self._tenant_sem_schema()

        with self.assertRaises(CommandError) as ctx:
            self._rodar(all=True)

        self.assertIn('1 tenant(s) com erro', str(ctx.exception))
        self.assertFalse(any(self._ativo(m) for m in self.legados.values()))


class ModeloPredefinidoTests(StorageTemporarioMixin, TenantTestCase):

    def setUp(self):
        super().setUp()
        User = get_user_model()
        User.objects.create_user(username='tester', password='senha123')
        self.client.login(username='tester', password='senha123')
        self.a = self._docx('Contrato A', 'contrato')
        self.b = self._docx('Contrato B', 'contrato')
        self.recibo = self._docx('Recibo A', 'recibo')

    def _docx(self, titulo, tipo, **kwargs):
        modelo = ModeloDocumento(titulo=titulo, tipo=tipo, **kwargs)
        modelo.arquivo.save('m.docx', ContentFile(b'x'), save=False)
        modelo.save()
        return modelo

    def _definir(self, modelo, logado=True):
        if not logado:
            self.client.logout()
        return self.client.post(
            reverse('documentos:definir_predefinido', args=[modelo.pk]), HTTP_HOST=self.domain.domain,
        )

    def _predefinidos(self, tipo):
        return set(ModeloDocumento.objects.filter(tipo=tipo, predefinido=True).values_list('titulo', flat=True))

    def test_marca_e_mostra_mensagem_e_redireciona(self):
        resp = self._definir(self.a)

        self.assertRedirects(resp, reverse('documentos:lista_modelos'), fetch_redirect_response=False)
        self.assertEqual(self._predefinidos('contrato'), {'Contrato A'})
        msgs = [str(m) for m in get_messages(resp.wsgi_request)]
        self.assertTrue(any('Contrato A' in m for m in msgs))

    def test_trocar_mantem_um_so_por_tipo(self):
        self._definir(self.a)
        self._definir(self.recibo)
        self._definir(self.b)

        self.assertEqual(self._predefinidos('contrato'), {'Contrato B'})
        self.assertEqual(self._predefinidos('recibo'), {'Recibo A'})

    def test_marcar_o_ja_predefinido_e_idempotente(self):
        self._definir(self.a)
        self._definir(self.a)

        self.assertEqual(self._predefinidos('contrato'), {'Contrato A'})

    def test_modelo_sem_arquivo_e_recusado(self):
        html = ModeloDocumento.objects.create(titulo='HTML', tipo='contrato', conteudo_html='<p>x</p>')

        self._definir(html)

        self.assertFalse(ModeloDocumento.objects.get(pk=html.pk).predefinido)

    def test_modelo_inativo_e_recusado(self):
        ModeloDocumento.objects.filter(pk=self.b.pk).update(ativo=False)

        self._definir(self.b)

        self.assertFalse(ModeloDocumento.objects.get(pk=self.b.pk).predefinido)

    def test_anonimo_e_redirecionado_para_login(self):
        resp = self._definir(self.a, logado=False)

        self.assertEqual(resp.status_code, 302)
        self.assertIn('login', resp['Location'])
        self.assertFalse(ModeloDocumento.objects.get(pk=self.a.pk).predefinido)

    def test_get_nao_e_permitido(self):
        resp = self.client.get(reverse('documentos:definir_predefinido', args=[self.a.pk]), HTTP_HOST=self.domain.domain)

        self.assertEqual(resp.status_code, 405)

    def test_constraint_impede_dois_ativos_predefinidos_no_mesmo_tipo(self):
        ModeloDocumento.objects.filter(pk=self.a.pk).update(predefinido=True)

        with self.assertRaises(IntegrityError), transaction.atomic():
            ModeloDocumento.objects.filter(pk=self.b.pk).update(predefinido=True)

    def test_desativar_via_save_desmarca_predefinido(self):
        self._definir(self.a)
        modelo = ModeloDocumento.objects.get(pk=self.a.pk)

        modelo.ativo = False
        modelo.save(update_fields=['ativo'])

        self.assertEqual(self._predefinidos('contrato'), set())
        self._definir(self.b)
        self.assertEqual(self._predefinidos('contrato'), {'Contrato B'})

    def test_excluir_predefinido_deixa_tipo_sem_predefinido(self):
        self._definir(self.a)

        ModeloDocumento.objects.get(pk=self.a.pk).delete()

        self.assertEqual(self._predefinidos('contrato'), set())

    def test_lista_mostra_botao_e_selo(self):
        self._definir(self.a)

        html = self.client.get(reverse('documentos:lista_modelos'), HTTP_HOST=self.domain.domain).content.decode()

        self.assertEqual(html.count('Padrão da imobiliária'), 1)
        self.assertEqual(html.count('Definir como padrão'), 2)  # b e recibo; a já é o padrão
        self.assertIn(reverse('documentos:definir_predefinido', args=[self.b.pk]), html)
        self.assertNotIn(reverse('documentos:definir_predefinido', args=[self.a.pk]), html)

    def test_lista_nao_oferece_botao_para_modelo_sem_arquivo(self):
        html_modelo = ModeloDocumento.objects.create(titulo='HTML', tipo='outro', conteudo_html='<p>x</p>')

        html = self.client.get(reverse('documentos:lista_modelos'), HTTP_HOST=self.domain.domain).content.decode()

        self.assertNotIn(reverse('documentos:definir_predefinido', args=[html_modelo.pk]), html)


class PredefinidoNosModelosPadraoTests(StorageTemporarioMixin, TenantTestCase):

    def _predefinidos(self):
        return set(ModeloDocumento.objects.filter(predefinido=True, ativo=True).values_list('tipo', flat=True))

    def test_tenant_novo_marca_um_por_tipo(self):
        resultado = criar_modelos_padrao_docx()

        self.assertEqual(self._predefinidos(), {'contrato', 'distrato', 'recibo'})
        self.assertEqual(len(resultado['predefinidos']), 3)

    def test_nao_marca_quando_tipo_ja_tem_predefinido(self):
        custom = ModeloDocumento(titulo='Meu contrato', tipo='contrato', predefinido=True)
        custom.arquivo.save('c.docx', ContentFile(b'x'), save=False)
        custom.save()

        resultado = criar_modelos_padrao_docx()

        self.assertEqual(
            list(ModeloDocumento.objects.filter(tipo='contrato', predefinido=True)), [custom],
        )
        self.assertEqual(self._predefinidos(), {'contrato', 'distrato', 'recibo'})
        self.assertNotIn(MODELOS_PADRAO_DOCX['contrato']['titulo'], resultado['predefinidos'])

    def test_predefinido_inativo_nao_bloqueia(self):
        antigo = ModeloDocumento.objects.create(titulo='Antigo', tipo='contrato', ativo=False)
        ModeloDocumento.objects.filter(pk=antigo.pk).update(predefinido=True)

        criar_modelos_padrao_docx()

        self.assertTrue(ModeloDocumento.objects.get(titulo=MODELOS_PADRAO_DOCX['contrato']['titulo']).predefinido)

    def test_dry_run_conta_sem_gravar(self):
        resultado = criar_modelos_padrao_docx(dry_run=True)

        self.assertEqual(len(resultado['predefinidos']), 3)
        self.assertEqual(ModeloDocumento.objects.count(), 0)

    def test_segunda_execucao_nao_muda_nada(self):
        criar_modelos_padrao_docx()
        estado = list(ModeloDocumento.objects.order_by('pk').values_list('pk', 'predefinido', 'ativo'))

        resultado = criar_modelos_padrao_docx()

        self.assertEqual(list(ModeloDocumento.objects.order_by('pk').values_list('pk', 'predefinido', 'ativo')), estado)
        self.assertEqual(resultado['predefinidos'], [])

    def test_nao_remarca_se_o_cliente_trocou_o_predefinido(self):
        criar_modelos_padrao_docx()
        ModeloDocumento.objects.filter(tipo='contrato').update(predefinido=False)

        criar_modelos_padrao_docx()  # DOCX já existe: não é criação, não marca de novo

        self.assertEqual(self._predefinidos(), {'distrato', 'recibo'})

    def test_comando_backfill_reporta_predefinidos(self):
        saida = io.StringIO()
        call_command(
            'backfill_modelos_padrao_docx', schema=self.tenant.schema_name, dry_run=True,
            stdout=saida, stderr=io.StringIO(),
        )

        self.assertIn('predefinidos=3', saida.getvalue())
        self.assertEqual(ModeloDocumento.objects.count(), 0)


class DownloadModeloExemploTests(TenantTestCase):

    def setUp(self):
        User = get_user_model()
        User.objects.create_user(username='tester', password='senha123')

    def _get(self, tipo):
        return self.client.get(
            reverse('documentos:download_modelo_exemplo', args=[tipo]), HTTP_HOST=self.domain.domain,
        )

    def test_download_devolve_o_docx_do_pacote(self):
        self.client.login(username='tester', password='senha123')

        for tipo in MODELOS_PADRAO_DOCX:
            with self.subTest(tipo=tipo):
                resp = self._get(tipo)
                self.assertEqual(resp.status_code, 200)
                self.assertEqual(
                    resp['Content-Type'],
                    'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                )
                self.assertIn('attachment', resp['Content-Disposition'])
                self.assertEqual(b''.join(resp.streaming_content), caminho_modelo_padrao_docx(tipo).read_bytes())

    def test_tipo_invalido_retorna_404(self):
        self.client.login(username='tester', password='senha123')

        for tipo in ('outro', 'CONTRATO', '..', '..%2F..%2Fsettings', 'contrato.docx'):
            with self.subTest(tipo=tipo):
                self.assertEqual(self._get(tipo).status_code, 404)

    def test_anonimo_e_redirecionado_para_login(self):
        resp = self._get('contrato')

        self.assertEqual(resp.status_code, 302)
        self.assertIn('login', resp['Location'])

    def test_tela_de_variaveis_mostra_secao_de_exemplos_e_aviso(self):
        self.client.login(username='tester', password='senha123')

        resp = self.client.get(reverse('documentos:variaveis'), HTTP_HOST=self.domain.domain)

        html_resp = resp.content.decode()
        self.assertIn('Modelos de exemplo', html_resp)
        self.assertIn('Use Times New Roman ou Arial para que o PDF mantenha a paginação do Word', html_resp)
        for tipo in MODELOS_PADRAO_DOCX:
            self.assertIn(reverse('documentos:download_modelo_exemplo', args=[tipo]), html_resp)

    def test_modal_novo_modelo_tem_link_para_a_secao(self):
        self.client.login(username='tester', password='senha123')

        resp = self.client.get(reverse('documentos:lista_modelos'), HTTP_HOST=self.domain.domain)

        self.assertIn('#modelos-exemplo', resp.content.decode())


class FiadorTextoTests(StorageTemporarioMixin, TenantTestCase):

    def setUp(self):
        super().setUp()
        imovel = Imovel.objects.create(
            codigo='IM-0001', tipo='apartamento', cep='60000000',
            logradouro='Rua Teste', numero='100', bairro='Centro',
            cidade='Fortaleza', estado='CE',
        )
        self.inquilino = Inquilino.objects.create(
            tipo='pf', nome='Rodrigo Oliveira', cpf='02738306006',
            telefone='85999999999', email='pagador@email.com',
            logradouro='Rua Doutor Vargas', numero='150',
            cidade='Porto Alegre', estado='RS', cep='91250000',
        )
        self.contrato = Contrato.objects.create(
            imovel=imovel, inquilino=self.inquilino, numero='CT-0001',
            data_inicio=date(2026, 1, 10), data_fim=date(2026, 12, 10),
            dia_vencimento=10, valor_aluguel=Decimal('1500.00'),
        )

    def _fiador_texto(self, **campos):
        for chave, valor in campos.items():
            setattr(self.inquilino, chave, valor)
        self.inquilino.save()
        self.contrato.refresh_from_db()
        return construir_contexto(self.contrato)['inquilino']['fiador_texto']

    def test_fiador_completo(self):
        texto = self._fiador_texto(
            fiador_nome='João da Silva', fiador_cpf='111.222.333-44', fiador_telefone='85988887777',
        )

        self.assertEqual(texto, 'FIADOR: João da Silva, CPF 111.222.333-44, telefone 85988887777.')

    def test_fiador_so_com_nome(self):
        self.assertEqual(self._fiador_texto(fiador_nome='João da Silva'), 'FIADOR: João da Silva.')

    def test_fiador_sem_cpf_omite_so_o_cpf(self):
        texto = self._fiador_texto(fiador_nome='João da Silva', fiador_telefone='85988887777')

        self.assertEqual(texto, 'FIADOR: João da Silva, telefone 85988887777.')

    def test_sem_fiador_string_vazia(self):
        self.assertEqual(self._fiador_texto(), '')

    def test_sem_nome_ignora_cpf_e_telefone(self):
        self.assertEqual(self._fiador_texto(fiador_cpf='111.222.333-44', fiador_telefone='85988887777'), '')

    def _texto_contrato_renderizado(self):
        criar_documentos_padrao()
        modelo = ModeloDocumento.objects.get(titulo=MODELOS_PADRAO_DOCX['contrato']['titulo'])
        saida = renderizar_docx(modelo, construir_contexto(self.contrato))
        return '\n'.join(p.text for p in Document(io.BytesIO(saida)).paragraphs)

    def test_contrato_sem_fiador_nao_deixa_resto(self):
        texto = self._texto_contrato_renderizado()

        self.assertNotIn('{{', texto)
        self.assertNotIn('FIADOR: ,', texto)
        self.assertNotIn('FIADOR', texto)

    def test_contrato_com_fiador_traz_a_clausula(self):
        self._fiador_texto(fiador_nome='João da Silva', fiador_cpf='111.222.333-44')

        texto = self._texto_contrato_renderizado()

        self.assertIn('FIADOR: João da Silva, CPF 111.222.333-44.', texto)
        self.assertNotIn('{{', texto)

    def test_recibo_nao_tem_mais_competencia(self):
        criar_documentos_padrao()
        modelo = ModeloDocumento.objects.get(titulo=MODELOS_PADRAO_DOCX['recibo']['titulo'])
        saida = renderizar_docx(modelo, construir_contexto(self.contrato))
        texto = '\n'.join(p.text for p in Document(io.BytesIO(saida)).paragraphs)

        self.assertNotIn('competência', texto)
        self.assertIn(', pago em ', texto)


class VariavelNovaEmTenantExistenteTests(StorageTemporarioMixin, TenantTestCase):
    """Tenant que já tem o catálogo antigo (sem fiador_texto) recebe só a variável nova."""

    def setUp(self):
        super().setUp()
        criar_documentos_padrao()
        VariavelDocumento.objects.filter(slug='inquilino.fiador_texto').delete()
        VariavelDocumento.objects.filter(slug='inquilino.nome').update(label='Rótulo do tenant')

    def test_criar_documentos_padrao_cria_so_a_variavel_faltante(self):
        resultado = criar_documentos_padrao()

        self.assertEqual(resultado['variaveis_criadas'], 1)
        nova = VariavelDocumento.objects.get(slug='inquilino.fiador_texto')
        self.assertEqual(nova.label, 'Texto do Fiador (cláusula completa)')
        self.assertEqual(nova.categoria, 'fiador')
        self.assertEqual(VariavelDocumento.objects.get(slug='inquilino.nome').label, 'Rótulo do tenant')

    def test_backfill_cria_a_variavel_faltante(self):
        saida = io.StringIO()
        call_command(
            'backfill_modelos_padrao_docx', schema=self.tenant.schema_name, stdout=saida, stderr=io.StringIO(),
        )

        self.assertIn('variáveis novas=1', saida.getvalue())
        self.assertTrue(VariavelDocumento.objects.filter(slug='inquilino.fiador_texto').exists())
        self.assertEqual(VariavelDocumento.objects.get(slug='inquilino.nome').label, 'Rótulo do tenant')

    def test_backfill_dry_run_so_conta(self):
        saida = io.StringIO()
        call_command(
            'backfill_modelos_padrao_docx', schema=self.tenant.schema_name, dry_run=True,
            stdout=saida, stderr=io.StringIO(),
        )

        self.assertIn('variáveis novas=1', saida.getvalue())
        self.assertFalse(VariavelDocumento.objects.filter(slug='inquilino.fiador_texto').exists())


class TaskGerarDocumentoSobFakeTenantTests(StorageTemporarioMixin, TenantTestCase):
    """O worker ativa o schema com schema_context: connection.tenant é um FakeTenant durante toda a geração."""

    def test_geracao_completa_sob_fake_tenant(self):
        contrato = criar_contrato_teste()
        modelo = criar_modelo_docx(montar_docx([['{{ contrato.numero }}']]))
        documento = ContratoDocumentoGerado.objects.create(
            contrato=contrato, modelo=modelo, titulo='Doc', status='pendente',
        )
        tipos_vistos = []

        def converter(docx_bytes):
            tipos_vistos.append(type(connection.tenant))
            return b'%PDF-fake'

        with patch('apps.documentos.tasks.converter_docx_para_pdf', side_effect=converter),                 patch('apps.documentos.tasks.sentry_sdk'):
            gerar_documento_docx(self.tenant.schema_name, str(documento.pk))

        documento.refresh_from_db()
        self.assertEqual(tipos_vistos, [FakeTenant])
        self.assertEqual(documento.status, 'gerado')
        self.assertTrue(documento.arquivo_pdf)
        self.assertIn(f'tenants/{self.tenant.schema_name}/documentos/contratos/', documento.arquivo_pdf.name)
